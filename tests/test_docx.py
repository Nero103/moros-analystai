from docx import Document

from document_utils import extract_docx_text


def test_extract_docx_text(tmp_path):
    file_path = tmp_path / "test.docx"

    document = Document()
    document.add_heading(
        "Moros AnalystAI Quarterly Report",
        level=1
    )
    document.add_paragraph(
        "Revenue increased 18 percent during the second quarter."
    )
    document.add_paragraph(
        "Customer retention improved from 82 percent to 89 percent."
    )
    document.add_paragraph(
        "The strongest growth occurred in the enterprise segment."
    )
    document.save(file_path)

    with open(file_path, "rb") as file:
        text = extract_docx_text(file)

    assert text is not None
    assert "Moros AnalystAI Quarterly Report" in text
    assert "Revenue increased 18 percent" in text
    assert "82 percent to 89 percent" in text
    assert "enterprise segment" in text

def test_empty_docx_returns_none(tmp_path):
    file_path = tmp_path / "empty.docx"

    document = Document()
    document.save(file_path)

    with open(file_path, "rb") as file:
        text = extract_docx_text(file)

    assert text is None

def test_invalid_docx_returns_none(tmp_path):
    file_path = tmp_path / "invalid.docx"

    file_path.write_text(
        "This is not a real Word document."
    )

    with open(file_path, "rb") as file:
        text = extract_docx_text(file)

    assert text is None